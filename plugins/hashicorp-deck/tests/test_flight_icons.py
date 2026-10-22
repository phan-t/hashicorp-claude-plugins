#!/usr/bin/env python3
"""Tests for the vendored Flight icons and the code that places them.

    python3 plugins/hashicorp-deck/tests/test_flight_icons.py

The helper tests need only the standard library. The rest need python-pptx and
beautifulsoup4, as the deck skill does, and are skipped without them.
"""

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
PLUGIN = HERE.parent
ICONS = PLUGIN / "assets" / "icons"
ASSETS = PLUGIN / "skills" / "hashicorp-deck-htmltopptx" / "assets"
TEMPLATE = PLUGIN / "skills" / "hashicorp-deck-html" / "assets" / "template.html"
sys.path.insert(0, str(ICONS))
sys.path.insert(0, str(ASSETS))

import flight  # noqa: E402

try:
    from pptx import Presentation
    import bs4  # noqa: F401
    import flight_icons
    import html_to_pptx
    from hashicorp_deck import Deck
except ImportError:
    Deck = None

needs_pptx = unittest.skipIf(Deck is None, "python-pptx and beautifulsoup4 not installed")


def icon_names(slide):
    return [s.name.split(": ", 1)[1] for s in slide.shapes if s.name.startswith("Icon: ")]


class Vendored(unittest.TestCase):
    def test_every_svg_is_in_the_catalog(self):
        svgs = {p.stem[:-3] for p in (ICONS / "flight").glob("*-24.svg")}
        self.assertEqual(svgs, {row[0] for row in flight.catalog()})

    def test_licence_ships_with_the_icons(self):
        self.assertIn("Mozilla Public License", (ICONS / "flight" / "LICENSE.md").read_text())


class Helper(unittest.TestCase):
    def test_search_matches_word_starts(self):
        names = [row[0] for row in flight.search("lock")]
        self.assertIn("lock", names)
        self.assertIn("cloud-lock", names)
        self.assertNotIn("clock", names)

    def test_ico_carries_the_name_for_the_converter(self):
        out = subprocess.run([sys.executable, str(ICONS / "flight.py"), "ico", "key"],
                             capture_output=True, text=True, check=True).stdout
        self.assertTrue(out.startswith('<div class="ico" data-icon="key"><svg aria-hidden="true"'))
        self.assertNotIn('width="24"', out)

    def test_unknown_name_suggests(self):
        r = subprocess.run([sys.executable, str(ICONS / "flight.py"), "ico", "usr-plus"],
                           capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("user-plus", r.stderr)


@needs_pptx
class Converter(unittest.TestCase):
    def segments(self, d):
        return flight_icons._parse(d)

    def test_every_icon_converts(self):
        for name in sorted(flight_icons.names()):
            with self.subTest(name=name):
                geom = flight_icons.custgeom(name)
                paths = geom.findall(".//{*}path")
                self.assertEqual(len(paths), len(flight_icons.svg_paths(name)))
                self.assertTrue(all(len(p) for p in paths))

    def test_glued_arc_flags(self):
        # "011.5" is large=0, sweep=1, x=1.5
        a = self.segments("M0 0a.75.75 0 011.5 0")
        b = self.segments("M0 0a.75 .75 0 0 1 1.5 0")
        self.assertEqual(a, b)
        self.assertAlmostEqual(a[-1][-2], 1.5)

    def test_drawing_after_close_gets_a_move(self):
        segs = self.segments("M1 1h2v2zc0 1 1 1 1 0")
        self.assertEqual(segs[4], ("M", 1.0, 1.0))
        self.assertEqual(segs[5][0], "C")

    def test_quadratic_becomes_cubic(self):
        segs = self.segments("M0 0q3 3 6 0")
        self.assertEqual(segs[1][0], "C")
        self.assertAlmostEqual(segs[1][1], 2.0)
        self.assertAlmostEqual(segs[1][2], 2.0)

    def test_unknown_name_suggests(self):
        with self.assertRaises(KeyError) as e:
            flight_icons.path_for("usr-plus")
        self.assertIn("user-plus", str(e.exception))


@needs_pptx
class Columns(unittest.TestCase):
    ICONS = ["key", "user-plus", "eye", "server"]

    def test_fills_every_square_in_both_kits(self):
        for mode in ("light", "dark"):
            for n in (2, 3, 4):
                with self.subTest(mode=mode, n=n):
                    d = Deck(mode=mode)
                    slide = d.columns("T", ["x"] * n, headings=["H"] * n,
                                      icons=self.ICONS[:n])
                    self.assertEqual(icon_names(slide), self.ICONS[:n])
                    squares = d.icon_squares(slide)
                    self.assertEqual(len(squares), n)
                    for sq, ic in zip(squares, [s for s in slide.shapes
                                               if s.name.startswith("Icon: ")]):
                        self.assertEqual(sq.left + sq.width // 2, ic.left + ic.width // 2)
                        self.assertEqual(sq.top + sq.height // 2, ic.top + ic.height // 2)
                    self.assertEqual(d.warnings, [])

    def test_icons_alone_pick_the_icon_layout(self):
        d = Deck()
        slide = d.columns("T", ["a", "b", "c"], icons=self.ICONS[:3])
        self.assertTrue(slide.slide_layout.name.startswith("Icon Text Container - 3"))

    def test_unknown_icon_raises_before_adding_a_slide(self):
        d = Deck()
        with self.assertRaises(KeyError):
            d.columns("T", ["a", "b"], headings=["A", "B"], icons=["key", "nope-nope"])
        self.assertEqual(len(d.prs.slides), 0)

    def test_count_must_match(self):
        with self.assertRaises(ValueError):
            Deck().columns("T", ["a", "b", "c"], headings=["A", "B", "C"], icons=["key"])

    def test_dark_v2_has_no_icon_layout(self):
        with self.assertRaises(ValueError):
            Deck(mode="dark-v2").columns("T", ["a", "b"], icons=["key", "eye"])

    def test_empty_squares_warn(self):
        d = Deck()
        d.columns("T", ["a", "b"], headings=["A", "B"])
        self.assertEqual(len(d.warnings), 1)
        self.assertIn("2 of 2 icon square(s) left empty", d.warnings[0])


@needs_pptx
class HtmlToPptx(unittest.TestCase):
    def convert(self, html):
        with tempfile.TemporaryDirectory() as tmp:
            src, out = os.path.join(tmp, "d.html"), os.path.join(tmp, "d.pptx")
            Path(src).write_text(html, encoding="utf-8")
            log = []
            html_to_pptx.convert(src, out, report=log.append)
            return Presentation(out), "\n".join(log)

    def test_template_cards_keep_their_icons(self):
        prs, log = self.convert(TEMPLATE.read_text(encoding="utf-8"))
        found = [icon_names(s) for s in prs.slides if icon_names(s)]
        self.assertEqual(found, [["key", "user-plus", "eye"]])
        self.assertIn("icons", log)

    def test_unknown_icon_is_reported_not_fatal(self):
        cards = "".join(f'<div class="card"><div class="ico" data-icon="{n}"></div>'
                        f'<h3>H{n}</h3><p class="desc">d</p></div>' for n in ("key", "nope-nope"))
        html = ('<section class="slide"><h1>Cover</h1></section>'
                f'<section class="slide"><h2>Cards</h2><div class="cards">{cards}</div></section>')
        prs, log = self.convert(html)
        self.assertIn("icons left out", log)
        self.assertEqual([icon_names(s) for s in prs.slides if icon_names(s)], [])

    def test_metric_cards_stay_metrics(self):
        cards = "".join(f'<div class="card"><div class="num">{v}</div><div class="cap">c</div></div>'
                        for v in ("92%", "3x"))
        html = ('<section class="slide"><h1>Cover</h1></section>'
                f'<section class="slide"><h2>Numbers</h2><div class="cards">{cards}</div></section>')
        _, log = self.convert(html)
        self.assertIn("metrics", log)


if __name__ == "__main__":
    unittest.main()
